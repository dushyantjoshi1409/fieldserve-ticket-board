#!/usr/bin/env python3
"""
FieldServe Ticket Board - Two-Way Sync Orchestrator
Architecture:
  1. Fetcher: Handles API pagination windowing (>500 tickets) & rate limiting (429 Retry-After)
  2. Parser & Extractor: Store Code (6-char alphanumeric), Work Type, Noise detection
  3. Workflow Engine: 6 sequential steps, coordinator id-parity, manager, engineer, finance
  4. 3-Way Diff & Conflict Resolver: Base (last sync) vs Local (sheet) vs Server (API)
  5. Pusher: PUT /api/v2/tickets/{id} with engineer name->id mapping and status str->int mapping
  6. Sheet Generator: Writes tickets.xlsx + Summary sheet with Action Owner workload counts
"""

import os
import sys
import time
import json
import re
from datetime import datetime, timezone, timedelta
from urllib.parse import urlparse
import requests

BASE_URL = os.environ.get("FIELDOPS_BASE_URL", "http://13.233.55.184:8377")
API_KEY = os.environ.get("FIELDOPS_API_KEY", "fk_90cf43b35e8eeaa60a0b6cf1")

SHEET_PATH = os.path.join(os.path.dirname(__file__), "tickets.xlsx")
STATE_PATH = os.path.join(os.path.dirname(__file__), "local_state.json")

STATUS_NUM_TO_STR = {2: "Open", 3: "Pending", 4: "Resolved", 5: "Closed"}
STATUS_STR_TO_NUM = {v.lower(): k for k, v in STATUS_NUM_TO_STR.items()}

WRITABLE_FIELDS = [
    "Status",
    "Engineer",
    "Manager",
    "Purchase Order",
    "Delivery Note",
    "Work Report",
    "Invoice",
]


# ==========================================
# 1. Date & URL Utilities
# ==========================================
def format_iso_clean(dt: datetime) -> str:
    """Server strictly requires YYYY-MM-DDTHH:MM:SSZ (no microseconds)."""
    if dt.tzinfo is None:
        dt = dt.replace(tzinfo=timezone.utc)
    else:
        dt = dt.astimezone(timezone.utc)
    return dt.strftime("%Y-%m-%dT%H:%M:%SZ")


def parse_iso_clean(s: str) -> datetime:
    """Parses ISO timestamp strings reliably."""
    if not s:
        return datetime(1970, 1, 1, tzinfo=timezone.utc)
    s = s.strip()
    # Strip microseconds if present
    s = re.sub(r"\.\d+", "", s)
    s = s.replace("Z", "+00:00")
    try:
        return datetime.fromisoformat(s).astimezone(timezone.utc)
    except Exception:
        return datetime(1970, 1, 1, tzinfo=timezone.utc)


def is_valid_url(v) -> bool:
    """Validates that a string is a real HTTP/HTTPS URL with valid netloc."""
    if not v or not isinstance(v, str):
        return False
    v = v.strip()
    if v.upper() == "NA":
        return False
    try:
        p = urlparse(v)
        return (
            p.scheme in ("http", "https")
            and bool(p.netloc)
            and "." in p.netloc
            and len(p.netloc) >= 3
        )
    except Exception:
        return False


def is_doc_met(value, waivable: bool = False) -> bool:
    """
    Returns True if document requirement is satisfied:
      - Valid URL
      - Or exact string 'NA' ONLY IF waivable is True.
    Junk text ('done', 'asked Ravi', 'ok', etc.) returns False.
    """
    if value is None:
        return False
    s = str(value).strip()
    if not s:
        return False
    if waivable and s == "NA":
        return True
    return is_valid_url(s)


# ==========================================
# 2. Heuristics: Store Code, Work Type, Noise
# ==========================================
def extract_store_code(subject: str = "", requester_email: str = "", requester_name: str = "") -> str:
    """
    Extracts 6-character alphanumeric store code (letters and digits).
    Looks in requester_email (store.k7q2xr@...), requester_name (OmniMart ...(K7Q2XR)),
    and subject (PM AT OMNIMART (K7Q2XR) or OMNIMART K7Q2XR).
    """
    common_words = {
        "CAMERA", "SYSTEM", "REPORT", "INVOICE", "BRANCH", "OFFICE",
        "STREET", "GROUND", "ALARMS", "SWITCH", "OUTLET", "STORES",
        "MARKET", "ONLINE", "VENDOR", "LABOUR", "REPAIR", "SERVER",
        "ROUTER", "SUPPLY", "BACKUP", "DESKTO", "SCREEN", "ACCESS",
        "SENSOR", "SECOND", "OMNIMA", "MARTVE", "DVRDVR"
    }

    # 1. Look in email: store.k7q2xr@omnimart.example
    if requester_email:
        m = re.search(r"store\.([a-zA-Z0-9]{6})@", str(requester_email).lower())
        if m:
            return m.group(1).upper()

    # 2. Look in requester_name: OmniMart Rarsi (M9GDCU)
    if requester_name:
        m = re.search(r"\(([A-Za-z0-9]{6})\)", str(requester_name))
        if m:
            return m.group(1).upper()

    # 3. Look in subject with parentheses: (K7Q2XR)
    if subject:
        subj = str(subject).upper()
        m = re.search(r"\(([A-Z0-9]{6})\)", subj)
        if m:
            code = m.group(1)
            if code not in common_words:
                return code

        # 4. Look for isolated tokens having both letter and digit (e.g. K7Q2XR, M9GDCU)
        tokens = re.findall(r"\b([A-Z0-9]{6})\b", subj)
        for token in tokens:
            if token in common_words:
                continue
            if re.search(r"[A-Z]", token) and re.search(r"[0-9]", token):
                return token

        # 5. If subject contains OMNIMART or STORE, look for 6-char token not in common words
        if "OMNIMART" in subj or "STORE" in subj:
            for token in tokens:
                if token not in common_words:
                    return token

    return ""


def extract_work_type(subject: str = "") -> str:
    """
    Classifies work type into:
      - Maintenance (PM / AMC)
      - Installation (INSTALL)
      - Service Call (faults, down, not working, complaints)
      - Other
    """
    s = f" {str(subject).upper()} "
    if " PM " in s or " AMC " in s or re.search(r"\bPM\b", s) or re.search(r"\bAMC\b", s):
        return "Maintenance"
    if "INSTALL" in s:
        return "Installation"
    service_keywords = [
        "NOT WORKING", "FAULT", "COMPLAINT", "DOWN",
        "ISSUE", "NOT SHOWING", "OFFLINE", "BLURRY",
        "REPAIR", "CHECK", "CAMERA", "FOOTAGE"
    ]
    if any(k in s for k in service_keywords):
        return "Service Call"
    return "Other"


def is_noise_ticket(subject: str = "", description: str = "") -> bool:
    """
    Detects noise tickets (spam, marketing, automated replies).
    Returns True if flagged as noise.
    """
    txt = f"{str(subject)} {str(description)}".lower()
    spam_keywords = [
        "offer", "listing", "business with", "sell smart", "earn more",
        "licenses", "best price", "discount", "promo", "lottery",
        "prize", "marketing", "newsletter", "loan", "seo", "domain",
        "investment", "automatic reply:", "out of office", "auto-reply",
        "autoreply", "undeliverable:"
    ]
    return any(k in txt for k in spam_keywords)


# ==========================================
# 3. Roster & Workflow Engine
# ==========================================
def build_roster_maps(agents: list):
    """
    Parses /api/v2/agents list and returns lookup structures.
    Coordinators are sorted alphabetically by name for deterministic id-parity assignment.
    """
    by_id = {}
    by_name = {}
    coordinators = []
    managers = []
    engineers = []
    finance = []

    for a in agents:
        by_id[a["id"]] = a
        by_name[a["name"]] = a
        role = a.get("role", "").strip().lower()
        if role == "coordinator":
            coordinators.append(a)
        elif role == "manager":
            managers.append(a)
        elif role == "engineer":
            engineers.append(a)
        elif role == "finance":
            finance.append(a)

    coordinators = sorted(coordinators, key=lambda x: x["name"])
    return {
        "by_id": by_id,
        "by_name": by_name,
        "coordinators": coordinators,
        "managers": managers,
        "engineers": engineers,
        "finance": finance,
    }


def compute_stage_and_owner(ticket: dict, roster: dict) -> tuple:
    """
    Calculates Stage and Action Owner based on §3 workflow:
      1. Assign Manager  -> Coordinator (sorted_by_name[ticket_id % 2])
      2. Assign Engineer -> Ticket's Manager
      3. Purchase Order  -> Ticket's Manager (valid URL only)
      4. Delivery Note   -> Finance (URL or 'NA')
      5. Work Report     -> Ticket's Engineer (valid URL only)
      6. Invoice         -> Finance (URL or 'NA')
      7. Done            -> Owner is blank ('')
    """
    tid = int(ticket.get("id", 0))
    manager_str = (ticket.get("manager") or "").strip()
    engineer_id = ticket.get("engineer_id")
    po = ticket.get("purchase_order") or ""
    dn = ticket.get("delivery_note") or ""
    wr = ticket.get("work_report") or ""
    inv = ticket.get("invoice") or ""

    coordinators = roster["coordinators"]
    finance = roster["finance"]
    finance_name = finance[0]["name"] if finance else "Finance"

    # Step 1: Assign Manager
    if not manager_str:
        if coordinators:
            coord_name = coordinators[tid % len(coordinators)]["name"]
        else:
            coord_name = "Coordinator"
        return "Assign Manager", coord_name

    # Step 2: Assign Engineer
    if not engineer_id:
        return "Assign Engineer", manager_str

    # Step 3: Purchase Order (Waivable: NO)
    if not is_doc_met(po, waivable=False):
        return "Purchase Order", manager_str

    # Step 4: Delivery Note (Waivable: YES)
    if not is_doc_met(dn, waivable=True):
        return "Delivery Note", finance_name

    # Step 5: Work Report (Waivable: NO)
    if not is_doc_met(wr, waivable=False):
        eng = roster["by_id"].get(engineer_id)
        eng_name = eng["name"] if eng else f"Engineer #{engineer_id}"
        return "Work Report", eng_name

    # Step 6: Invoice (Waivable: YES)
    if not is_doc_met(inv, waivable=True):
        return "Invoice", finance_name

    # Step 7: Done
    return "Done", ""


def compute_integrity_flags(ticket: dict, roster: dict) -> str:
    """
    Surfaces integrity problems:
      - Closed ticket with missing paperwork
      - Resolved ticket without Work Report
      - Unknown Manager not in roster
    """
    status = ticket.get("status")
    flags = []

    # Check 6 steps
    has_mgr = bool((ticket.get("manager") or "").strip())
    has_eng = bool(ticket.get("engineer_id"))
    po_met = is_doc_met(ticket.get("purchase_order"), waivable=False)
    dn_met = is_doc_met(ticket.get("delivery_note"), waivable=True)
    wr_met = is_doc_met(ticket.get("work_report"), waivable=False)
    inv_met = is_doc_met(ticket.get("invoice"), waivable=True)
    all_met = has_mgr and has_eng and po_met and dn_met and wr_met and inv_met

    # Closed (5) requires all steps met
    if status == 5 and not all_met:
        flags.append("Closed but paperwork missing")

    # Resolved (4) requires Work Report
    if status == 4 and not wr_met:
        flags.append("Resolved without Work Report")

    # Manager validation
    mgr_str = (ticket.get("manager") or "").strip()
    if mgr_str and mgr_str not in roster["by_name"]:
        flags.append("Manager not in roster")

    return "; ".join(flags)


# ==========================================
# 4. API Client (Rate Limit & Windowing)
# ==========================================
def api_auth():
    return (API_KEY, "x")


def fetch_agents():
    """Fetches agents list with rate-limiting retries."""
    url = f"{BASE_URL}/api/v2/agents"
    for attempt in range(5):
        r = requests.get(url, auth=api_auth(), timeout=15)
        if r.status_code == 429:
            retry_sec = int(r.headers.get("Retry-After", "2"))
            time.sleep(retry_sec + 1)
            continue
        r.raise_for_status()
        return r.json()
    raise RuntimeError("Failed to fetch agents after 5 retries")


def fetch_all_tickets_windowed(updated_since_str: str = "1970-01-01T00:00:00Z") -> dict:
    """
    Enumerates all tickets via pagination windowing:
      - API limits page <= 10 and per_page <= 50 (max 500 tickets per window).
      - After reaching page 10, advances updated_since to the last ticket's updated_at.
      - Automatically sleeps on HTTP 429 (Retry-After).
    """
    all_tickets = {}
    current_since = updated_since_str
    per_page = 50
    max_pages = 10

    while True:
        window_tickets = 0
        last_updated_at = None

        for page in range(1, max_pages + 1):
            url = f"{BASE_URL}/api/v2/tickets"
            params = {"page": page, "per_page": per_page, "updated_since": current_since}

            resp = None
            for _ in range(5):
                resp = requests.get(url, auth=api_auth(), params=params, timeout=20)
                if resp.status_code == 429:
                    retry_sec = int(resp.headers.get("Retry-After", "2"))
                    print(f"[Rate-Limit] HTTP 429, sleeping {retry_sec + 1}s...", file=sys.stderr)
                    time.sleep(retry_sec + 1)
                    continue
                if resp.status_code == 400 and page > 1:
                    # Page beyond data for current window
                    break
                resp.raise_for_status()
                break

            if not resp or resp.status_code != 200:
                break

            data = resp.json()
            tickets = data if isinstance(data, list) else data.get("tickets", [])
            if not tickets:
                break

            for t in tickets:
                all_tickets[t["id"]] = t
                last_updated_at = t.get("updated_at", last_updated_at)
                window_tickets += 1

            if len(tickets) < per_page:
                break

        if window_tickets == 0:
            break

        # Check if we hit the 500-ticket window cap and need to advance window
        if window_tickets >= (max_pages * per_page) and last_updated_at:
            if last_updated_at == current_since:
                # All 500 tickets have the exact same second, advance by 1 second
                dt = parse_iso_clean(last_updated_at) + timedelta(seconds=1)
                current_since = format_iso_clean(dt)
            else:
                current_since = last_updated_at
            continue
        else:
            # Under 500 tickets means we fetched all tickets in this window
            break

    return all_tickets


def put_ticket(ticket_id: int, payload: dict) -> tuple:
    """
    Sends partial update via PUT /api/v2/tickets/{id}.
    Returns (success: bool, response_json_or_error: str).
    """
    url = f"{BASE_URL}/api/v2/tickets/{ticket_id}"
    for _ in range(5):
        r = requests.put(url, auth=api_auth(), json=payload, timeout=15)
        if r.status_code == 429:
            retry_sec = int(r.headers.get("Retry-After", "2"))
            time.sleep(retry_sec + 1)
            continue
        if r.status_code in (400, 422):
            return False, f"HTTP {r.status_code}: {r.text}"
        r.raise_for_status()
        return True, r.json()
    return False, "Rate limit persistent"


# ==========================================
# 5. Spreadsheet I/O
# ==========================================
def load_sheet_tickets() -> dict:
    """Reads current tickets from tickets.xlsx into a dict keyed by Ticket ID."""
    import openpyxl

    if not os.path.exists(SHEET_PATH):
        return {}

    wb = openpyxl.load_workbook(SHEET_PATH, data_only=True)
    ws = wb.active
    headers = [c.value for c in ws[1]]
    if not headers or "Ticket ID" not in headers:
        return {}

    col_idx = {h: i for i, h in enumerate(headers)}
    sheet_tickets = {}

    for row in ws.iter_rows(min_row=2, values_only=True):
        tid_raw = row[col_idx["Ticket ID"]]
        if tid_raw is None or str(tid_raw).strip() == "":
            continue
        try:
            tid = int(str(tid_raw).strip())
        except ValueError:
            continue

        sheet_tickets[tid] = {
            "Ticket ID": tid,
            "Subject": row[col_idx["Subject"]] if "Subject" in col_idx else "",
            "Status": str(row[col_idx["Status"]]).strip() if "Status" in col_idx and row[col_idx["Status"]] is not None else "",
            "Engineer": str(row[col_idx["Engineer"]]).strip() if "Engineer" in col_idx and row[col_idx["Engineer"]] is not None else "",
            "Manager": str(row[col_idx["Manager"]]).strip() if "Manager" in col_idx and row[col_idx["Manager"]] is not None else "",
            "Purchase Order": str(row[col_idx["Purchase Order"]]).strip() if "Purchase Order" in col_idx and row[col_idx["Purchase Order"]] is not None else "",
            "Delivery Note": str(row[col_idx["Delivery Note"]]).strip() if "Delivery Note" in col_idx and row[col_idx["Delivery Note"]] is not None else "",
            "Work Report": str(row[col_idx["Work Report"]]).strip() if "Work Report" in col_idx and row[col_idx["Work Report"]] is not None else "",
            "Invoice": str(row[col_idx["Invoice"]]).strip() if "Invoice" in col_idx and row[col_idx["Invoice"]] is not None else "",
            "Stage": str(row[col_idx["Stage"]]).strip() if "Stage" in col_idx and row[col_idx["Stage"]] is not None else "",
            "Action Owner": str(row[col_idx["Action Owner"]]).strip() if "Action Owner" in col_idx and row[col_idx["Action Owner"]] is not None else "",
        }

    return sheet_tickets


def write_sheet_tickets(all_tickets: dict, roster: dict):
    """
    Writes all tickets to tickets.xlsx with required columns + computed fields,
    and creates a Summary sheet showing tickets waiting per Action Owner.
    """
    import openpyxl
    from collections import Counter

    wb = openpyxl.Workbook()

    # Main Tickets sheet
    ws = wb.active
    ws.title = "Tickets"

    headers = [
        "Ticket ID", "Subject", "Status", "Engineer", "Manager",
        "Purchase Order", "Delivery Note", "Work Report", "Invoice",
        "Stage", "Action Owner", "Store Code", "Work Type", "Is Noise",
        "Integrity Flag", "Updated At"
    ]
    ws.append(headers)

    owner_workload = Counter()

    for tid, t in sorted(all_tickets.items()):
        eng_id = t.get("engineer_id")
        eng_name = ""
        if eng_id and eng_id in roster["by_id"]:
            eng_name = roster["by_id"][eng_id]["name"]

        status_str = STATUS_NUM_TO_STR.get(t.get("status"), str(t.get("status", "Open")))
        stage, owner = compute_stage_and_owner(t, roster)
        flags = compute_integrity_flags(t, roster)
        store_code = extract_store_code(t.get("subject", ""), t.get("requester_email", ""), t.get("requester_name", ""))
        work_type = extract_work_type(t.get("subject", ""))
        is_noise = is_noise_ticket(t.get("subject", ""), t.get("description", ""))

        if owner and stage != "Done":
            owner_workload[owner] += 1

        ws.append([
            tid,
            t.get("subject", "") or "",
            status_str,
            eng_name,
            t.get("manager", "") or "",
            t.get("purchase_order", "") or "",
            t.get("delivery_note", "") or "",
            t.get("work_report", "") or "",
            t.get("invoice", "") or "",
            stage,
            owner,
            store_code,
            work_type,
            "YES" if is_noise else "NO",
            flags,
            t.get("updated_at", "") or ""
        ])

    # Summary Sheet: Action Owner Workload
    ws_summary = wb.create_sheet(title="Summary")
    ws_summary.append(["Action Owner", "Tickets Waiting"])
    for owner, count in owner_workload.most_common():
        ws_summary.append([owner, count])

    wb.save(SHEET_PATH)
    print(f"[Sheet] Successfully written {len(all_tickets)} tickets to {SHEET_PATH}")


# ==========================================
# 6. Two-Way Sync Orchestrator
# ==========================================
def run_sync():
    """
    Executes one complete two-way sync cycle:
      1. Load local base state snapshot from local_state.json.
      2. Fetch agents roster from /api/v2/agents.
      3. Incremental or full fetch of tickets from server.
      4. Load sheet tickets from tickets.xlsx.
      5. 3-Way diff (Base vs Sheet vs Server) on writable fields.
      6. Push sheet edits to API via PUT (with engineer & status mapping).
      7. Rewrite tickets.xlsx with latest synchronized state.
      8. Persist new snapshot in local_state.json.
    """
    print(f"=== FieldServe Sync Start [BASE={BASE_URL}] ===")

    # 1. Load Local State Snapshot
    state = {"last_sync": "1970-01-01T00:00:00Z", "tickets": {}, "agents": []}
    if os.path.exists(STATE_PATH):
        try:
            with open(STATE_PATH, "r", encoding="utf-8") as f:
                state = json.load(f)
        except Exception as e:
            print(f"[Warn] Could not read {STATE_PATH}: {e}, using clean state")

    # 2. Fetch Agents Roster
    try:
        agents = fetch_agents()
    except Exception as e:
        print(f"[Warn] API agents fetch failed ({e}). Attempting cached agents...")
        agents = state.get("agents", [])

    if not agents:
        print("[Error] No agents available from API or cache. Aborting sync.", file=sys.stderr)
        sys.exit(1)

    roster = build_roster_maps(agents)

    # 3. Fetch Server Tickets
    base_tickets = state.get("tickets", {})
    last_sync_str = state.get("last_sync", "1970-01-01T00:00:00Z")

    # Determine fetch strategy:
    # If base cache is empty or sheet missing, do full fetch from epoch
    if not base_tickets or not os.path.exists(SHEET_PATH):
        fetch_since = "1970-01-01T00:00:00Z"
        print(f"[Sync] Initializing: Performing full fetch from epoch...")
    else:
        # Incremental fetch with safety buffer of 60 seconds for server clock skew
        last_dt = parse_iso_clean(last_sync_str)
        buffered_dt = max(datetime(1970, 1, 1, tzinfo=timezone.utc), last_dt - timedelta(seconds=60))
        fetch_since = format_iso_clean(buffered_dt)
        print(f"[Sync] Incremental fetch with window since {fetch_since}...")

    server_fetched = fetch_all_tickets_windowed(fetch_since)
    print(f"[API] Fetched {len(server_fetched)} tickets from server")

    # Reconstruct server tickets map from cached raw tickets + fresh updates
    server_tickets = {}
    for tid_str, entry in base_tickets.items():
        if "raw" in entry:
            server_tickets[int(tid_str)] = entry["raw"]

    for tid, t in server_fetched.items():
        server_tickets[tid] = t

    # 4. Load Sheet Tickets
    sheet_tickets = load_sheet_tickets()
    print(f"[Sheet] Loaded {len(sheet_tickets)} tickets from spreadsheet")

    # 5. 3-Way Diff & Conflict Resolution
    conflicts_log = []
    pushed_count = 0

    # Build base comparable dict: tid -> dict of field -> str
    base_comp_map = {}
    for tid_str, entry in base_tickets.items():
        try:
            tid = int(tid_str)
            base_comp_map[tid] = entry.get("comp", {})
        except ValueError:
            pass

    # Check for edits across all known ticket IDs
    all_tids = set(server_tickets.keys()) | set(sheet_tickets.keys())

    for tid in sorted(all_tids):
        server_raw = server_tickets.get(tid)
        sheet_row = sheet_tickets.get(tid)
        base_comp = base_comp_map.get(tid)

        # Ticket only exists on one side
        if not server_raw or not sheet_row:
            continue

        # Map current server fields to string representations
        eng_id = server_raw.get("engineer_id")
        eng_name = roster["by_id"][eng_id]["name"] if eng_id and eng_id in roster["by_id"] else ""
        server_comp = {
            "Status": STATUS_NUM_TO_STR.get(server_raw.get("status"), str(server_raw.get("status", ""))),
            "Engineer": eng_name,
            "Manager": (server_raw.get("manager") or "").strip(),
            "Purchase Order": (server_raw.get("purchase_order") or "").strip(),
            "Delivery Note": (server_raw.get("delivery_note") or "").strip(),
            "Work Report": (server_raw.get("work_report") or "").strip(),
            "Invoice": (server_raw.get("invoice") or "").strip(),
        }

        # If base is missing for this ticket (first time seen), initialize base
        if not base_comp:
            base_comp = dict(server_comp)

        # Compute field-level diffs
        payload = {}
        for field in WRITABLE_FIELDS:
            sheet_val = str(sheet_row.get(field, "")).strip()
            base_val = str(base_comp.get(field, "")).strip()
            server_val = str(server_comp.get(field, "")).strip()

            local_changed = (sheet_val != base_val)
            server_changed = (server_val != base_val)

            if local_changed and not server_changed:
                # Local sheet was edited -> Prepare payload for PUT
                if field == "Status":
                    status_int = STATUS_STR_TO_NUM.get(sheet_val.lower())
                    if status_int:
                        payload["status"] = status_int
                elif field == "Engineer":
                    if not sheet_val:
                        payload["engineer_id"] = None
                    else:
                        eng_match = roster["by_name"].get(sheet_val)
                        if eng_match:
                            payload["engineer_id"] = eng_match["id"]
                        else:
                            print(f"[Warn] Unknown engineer '{sheet_val}' for ticket {tid}, skipping field")
                elif field == "Manager":
                    payload["manager"] = sheet_val
                elif field == "Purchase Order":
                    payload["purchase_order"] = sheet_val
                elif field == "Delivery Note":
                    payload["delivery_note"] = sheet_val
                elif field == "Work Report":
                    payload["work_report"] = sheet_val
                elif field == "Invoice":
                    payload["invoice"] = sheet_val

            elif local_changed and server_changed:
                if sheet_val == server_val:
                    # Both changed to identical value -> no conflict
                    pass
                else:
                    # Conflict: both changed differently
                    # Documented Policy: Server Wins
                    conflicts_log.append({
                        "ticket_id": tid,
                        "field": field,
                        "base": base_val,
                        "local": sheet_val,
                        "server": server_val,
                        "resolution": "server_wins"
                    })
                    print(f"[Conflict] Ticket {tid} Field '{field}': Local='{sheet_val}' vs Server='{server_val}'. Server wins.")

            # If unchanged on sheet (local_changed is False), sheet never overwrites server.

        # If we have local changes to push
        if payload:
            ok, resp = put_ticket(tid, payload)
            if ok:
                pushed_count += 1
                print(f"[Push] Ticket {tid} updated on server: {payload}")
                # Update server_tickets in-memory
                for k, v in payload.items():
                    server_raw[k] = v
                if isinstance(resp, dict) and "updated_at" in resp:
                    server_raw["updated_at"] = resp["updated_at"]
            else:
                print(f"[Error] Failed to PUT ticket {tid} with payload {payload}: {resp}", file=sys.stderr)

    # 6. Rewrite Sheet with Latest Server State
    write_sheet_tickets(server_tickets, roster)

    # 7. Update and Save Local State Snapshot
    new_snapshot_tickets = {}
    for tid, t in server_tickets.items():
        eng_id = t.get("engineer_id")
        eng_name = roster["by_id"][eng_id]["name"] if eng_id and eng_id in roster["by_id"] else ""
        new_snapshot_tickets[str(tid)] = {
            "updated_at": t.get("updated_at", ""),
            "comp": {
                "Status": STATUS_NUM_TO_STR.get(t.get("status"), str(t.get("status", ""))),
                "Engineer": eng_name,
                "Manager": (t.get("manager") or "").strip(),
                "Purchase Order": (t.get("purchase_order") or "").strip(),
                "Delivery Note": (t.get("delivery_note") or "").strip(),
                "Work Report": (t.get("work_report") or "").strip(),
                "Invoice": (t.get("invoice") or "").strip(),
            },
            "raw": t
        }

    now_iso = format_iso_clean(datetime.now(timezone.utc))
    state_out = {
        "last_sync": now_iso,
        "tickets": new_snapshot_tickets,
        "agents": agents,
        "conflicts": conflicts_log
    }

    with open(STATE_PATH, "w", encoding="utf-8") as f:
        json.dump(state_out, f, indent=2)

    print(f"=== Sync Completed: {pushed_count} pushed, {len(conflicts_log)} conflicts, Total {len(server_tickets)} tickets ===")


if __name__ == "__main__":
    run_sync()
